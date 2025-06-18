import time
import allure
import os
from selenium.webdriver.common.by import By

from util_tools.basePage import BasePage


class UserImportPage(BasePage):
    # 定义页面URL地址
    url = '#/admin/system/user/index'

    # 定位"导入"按钮的定位器，按钮具有类名'el-button--primary'，并且包含文本'span'为"导入"
    file_input_click = (By.XPATH, "//button[contains(@class, 'el-button--primary')][contains(span, '导入')]")
    # 定位文件上传的input元素，通过CSS选择器找到具有类名'el-upload__input'的元素
    file_input_locator = (By.CSS_SELECTOR, "input.el-upload__input")
    # 定位"确认"按钮的定位器，使用XPath找到文本为"确认"的按钮
    submit_import = (By.XPATH, "//button[span[text()='确认']]")

    def user_import(self):
        """上传Excel文件的操作流程"""
        self.open_url(self.url)  # 打开用户管理页面
        allure.attach(self.url, '打开测试页面', attachment_type=allure.attachment_type.TEXT)  # 添加操作日志到报告中
        time.sleep(1)  # 等待页面加载
        self.click(self.file_input_click)  # 点击"导入"按钮，打开文件上传对话框
        time.sleep(1)  # 等待文件上传对话框弹出
        
        # 使用项目相对路径，指向Excel文件
        file_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../testcase/test_data/user_import_template.xlsx'))
        
        # 检查Excel文件是否存在，如果不存在则创建
        if not os.path.exists(file_path):
            self._create_excel_template(file_path)
        
        allure.attach(f"使用Excel文件路径: {file_path}", '文件路径信息', attachment_type=allure.attachment_type.TEXT)
        
        self.send_keys(self.file_input_locator, file_path)  # 在文件上传input元素中输入文件路径，模拟文件上传
        time.sleep(1)  # 等待文件上传完成
        self.click(self.submit_import)  # 点击"确认"按钮提交文件
        time.sleep(1)  # 等待文件上传和导入操作完成

    def _create_excel_template(self, file_path):
        """创建Excel模板文件"""
        try:
            from openpyxl import Workbook
            
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            ws.title = "用户导入模板"
            
            # 设置表头
            headers = [
                "用户名", "姓名", "昵称", "手机号", "邮箱", 
                "部门名称", "岗位名称", "角色", "锁定标记"
            ]
            
            # 写入表头
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            # 添加测试数据
            test_data = [
                ["testuser1", "测试用户1", "测试1", "13800138001", "testuser1@test.com", "技术部", "产品经理", "产品经理", 0],
                ["testuser2", "测试用户2", "测试2", "13800138002", "testuser2@test.com", "技术部", "部门总监", "部门总监", 0],
                ["testuser3", "测试用户3", "测试3", "13800138003", "testuser3@test.com", "市场部", "普通员工", "普通用户", 0]
            ]
            
            # 写入测试数据
            for row_idx, row_data in enumerate(test_data, 2):
                for col_idx, cell_value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
            # 确保目录存在
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            
            # 保存文件
            wb.save(file_path)
            
            allure.attach(f"已自动创建Excel模板文件: {file_path}", '文件创建', attachment_type=allure.attachment_type.TEXT)
            
        except ImportError:
            allure.attach("缺少openpyxl库，无法创建Excel文件", '错误信息', attachment_type=allure.attachment_type.TEXT)
            raise Exception("请安装openpyxl库: pip install openpyxl")
        except Exception as e:
            allure.attach(f"创建Excel文件失败: {str(e)}", '错误信息', attachment_type=allure.attachment_type.TEXT)
            raise

